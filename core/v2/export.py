"""V2 测试用例导出 - 从结构化 TestCase 生成 Excel / Markdown。

与 V1 的 core/output.py 解耦：V2 以 Pydantic TestCase 为输入，保留状态/来源/测试数据等结构化信息。
默认由调用方过滤 ARCHIVED 用例；本模块只负责格式化。
"""

from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.schemas import TestCase

HEADERS = [
    "用例编号",
    "模块",
    "标题",
    "前置条件",
    "测试步骤",
    "预期结果",
    "测试数据",
    "优先级",
    "用例类型",
    "状态",
    "来源",
    "生成方式",
    "备注",
]

PRIORITY_COLORS = {
    "P0": "FF0000",
    "P1": "FF6600",
    "P2": "FFCC00",
    "P3": "00CC00",
}

STATUS_LABELS = {
    "generated": "已生成",
    "validated": "已校验",
    "validation_failed": "校验失败",
    "reviewed": "已评审",
    "edited": "已编辑",
    "re_review_required": "待重新评审",
    "confirmed": "已确认",
    "archived": "已归档",
}

PROVENANCE_LABELS = {
    "llm": "AI 生成",
    "strategy": "策略引擎",
    "human": "人工",
    "validator": "校验器",
    "optimizer": "优化器",
    "migrated": "V1 迁移",
    "xmind": "XMind",
}

GENERATION_MODE_LABELS = {
    "code": "代码模板",
    "llm": "LLM 合成",
    "hybrid": "混合",
}

_INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n]+')


def _enum_value(value) -> str:
    return value.value if hasattr(value, "value") else str(value)


def _label(mapping: dict[str, str], value) -> str:
    raw = _enum_value(value)
    return mapping.get(raw, raw)


def _format_data_plan(case: TestCase) -> str:
    parts = []
    for item in case.data_plan:
        value = "" if item.value is None else str(item.value)
        meta = "/".join(x for x in (item.strategy, item.source) if x)
        parts.append(f"{item.field}={value}" + (f" ({meta})" if meta else ""))
    return "\n".join(parts)


def _case_row(case: TestCase) -> list[str]:
    return [
        case.display_id,
        case.module,
        case.title,
        case.precondition,
        case.render_steps_text(),
        case.expected,
        _format_data_plan(case),
        _enum_value(case.priority),
        _enum_value(case.type),
        _label(STATUS_LABELS, case.status),
        _label(PROVENANCE_LABELS, case.provenance),
        _label(GENERATION_MODE_LABELS, case.generation_mode),
        case.remark,
    ]


def _write_summary_sheet(
    ws,
    *,
    run_title: str,
    run_id: str,
    exported_at: datetime,
    case_count: int,
    archived_excluded: int,
) -> None:
    ws.title = "运行信息"
    ws.append(["字段", "值"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws.append(["需求标题", run_title])
    ws.append(["Run ID", run_id])
    ws.append(["导出时间", exported_at.strftime("%Y-%m-%d %H:%M:%S")])
    ws.append(["导出用例数", case_count])
    ws.append(["已排除归档用例", archived_excluded])
    ws.append(["说明", "默认不导出已归档用例；测试数据来自结构化 DataPlan。"])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 60
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_test_cases_excel(
    cases: list[TestCase],
    *,
    run_title: str,
    run_id: str,
    archived_excluded: int = 0,
    exported_at: datetime | None = None,
) -> bytes:
    """生成 Excel 二进制内容（两个 sheet：测试用例 + 运行信息）。"""
    exported_at = exported_at or datetime.now()
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for row_idx, case in enumerate(cases, 2):
        for col_idx, value in enumerate(_case_row(case), 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if col_idx == 8:  # 优先级列
                color = PRIORITY_COLORS.get(str(value), "")
                if color:
                    cell.font = Font(bold=True, color=color)

    widths = [12, 15, 30, 22, 44, 30, 30, 10, 14, 14, 12, 12, 20]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(cases) + 1}"

    _write_summary_sheet(
        wb.create_sheet(),
        run_title=run_title,
        run_id=run_id,
        exported_at=exported_at,
        case_count=len(cases),
        archived_excluded=archived_excluded,
    )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _md_cell(text: str) -> str:
    return str(text).replace("|", r"\|").replace("\r\n", "\n").replace("\n", "<br>")


def export_test_cases_markdown(
    cases: list[TestCase],
    *,
    run_title: str,
    run_id: str,
    archived_excluded: int = 0,
    exported_at: datetime | None = None,
) -> str:
    """生成 Markdown 文本（按模块分组；默认调用方已排除 ARCHIVED）。"""
    exported_at = exported_at or datetime.now()
    lines = [
        f"# {run_title} - 测试用例",
        "",
        f"> Run ID: {run_id}",
        f"> 导出时间: {exported_at.strftime('%Y-%m-%d %H:%M:%S')}",
        f"> 用例总数: {len(cases)}",
    ]
    if archived_excluded:
        lines.append(f"> 已排除归档用例: {archived_excluded}")
    lines.append("")

    modules: dict[str, list[TestCase]] = {}
    for case in cases:
        modules.setdefault(case.module or "未分类", []).append(case)

    lines.extend(["## 用例总览", "", "| 模块 | 用例数 |", "|------|-------:|"])
    for module, grouped in modules.items():
        lines.append(f"| {_md_cell(module)} | {len(grouped)} |")
    lines.append("")

    for module, grouped in modules.items():
        lines.extend([f"## {module}", ""])
        for case in grouped:
            lines.extend(
                [
                    f"### {case.display_id} - {case.title}",
                    "",
                    f"- **优先级**: {_enum_value(case.priority)}",
                    f"- **类型**: {_enum_value(case.type)}",
                    f"- **状态**: {_label(STATUS_LABELS, case.status)}",
                    f"- **来源**: {_label(PROVENANCE_LABELS, case.provenance)}",
                    f"- **生成方式**: {_label(GENERATION_MODE_LABELS, case.generation_mode)}",
                ]
            )
            if case.precondition:
                lines.append(f"- **前置条件**: {case.precondition}")
            data_plan = _format_data_plan(case)
            if data_plan:
                lines.append("- **测试数据**:")
                lines.extend(f"  - {line}" for line in data_plan.splitlines())
            lines.append("- **测试步骤**:")
            for step in case.render_steps_text().splitlines():
                if step.strip():
                    lines.append(f"  {step}")
            lines.append(f"- **预期结果**: {case.expected}")
            if case.remark:
                lines.append(f"- **备注**: {case.remark}")
            lines.append("")

    return "\n".join(lines)


def safe_export_filename(
    run_title: str, *, ext: str, exported_at: datetime | None = None, label: str = "测试用例"
) -> str:
    """生成安全的下载文件名（保留中文，去掉路径/控制字符）。"""
    exported_at = exported_at or datetime.now()
    base = _INVALID_FILENAME_CHARS.sub("_", (run_title or "").strip()).strip(" .") or "testcases"
    return f"{base[:60]}_{label}_{exported_at.strftime('%Y%m%d_%H%M%S')}.{ext}"
