"""Step 12.1 验收测试 - V2 导出（Excel / Markdown）。

覆盖：导出内容、默认排除 ARCHIVED、格式校验、404/401/503、Run 归属校验、空运行、文件名安全。
全离线：不调用 LLM、不写生产库（使用临时 V2 库）。
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from core.schemas import (
    DataPlanItem,
    GenerationConfig,
    GenerationMode,
    Priority,
    Provenance,
    RequirementDoc,
    RequirementVersion,
    Run,
    RunStatus,
    SourceType,
    TestCase,
    TestCaseStatus,
    TestCaseType,
    TestStep,
    new_ulid,
)
from core.v2 import export as export_lib
from core.v2 import repository as repo


@pytest.fixture
def v2_db_path(tmp_path):
    from core.v2 import db as v2db
    from core.v2 import ddl

    old = v2db.get_v2_db_path()
    path = str(tmp_path / "export_v2.db")
    v2db.set_v2_db_path(path)
    ddl.create_v2_schema()
    yield path
    v2db.set_v2_db_path(old)


@pytest.fixture
def v2_client(auth_client, v2_db_path):
    return auth_client


@pytest.fixture
def v2_anon(client, v2_db_path):
    return client


def _session_v2_uid(client) -> str:
    legacy_uid = client.get("/api/me").get_json()["user"]["id"]
    row = repo.get_user_by_legacy_id(legacy_uid)
    if row is not None:
        return row["id"]
    uid = new_ulid()
    repo.save_user(uid, "export_user", "", legacy_int_id=legacy_uid)
    return uid


def _create_run(uid: str, title: str = "退款需求") -> Run:
    doc = RequirementDoc(user_id=uid, title=title, source_type=SourceType.TEXT)
    repo.save_doc(doc)
    version = RequirementVersion(doc_id=doc.id, version_no=1, raw_text="退款需求", provenance=Provenance.LLM)
    repo.save_version(version)
    cfg = GenerationConfig(
        model_provider="test", model_name="fake", temperature=0.3, prompt_version="p", generator_version="g"
    )
    repo.save_generation_config(cfg)
    run = Run(
        user_id=uid,
        doc_id=doc.id,
        requirement_version_id=version.id,
        generation_config_id=cfg.id,
        status=RunStatus.DONE,
    )
    repo.save_run(run)
    return run


def _case(
    run: Run,
    *,
    tc_id: str,
    display_id: str,
    status: TestCaseStatus,
    module: str = "登录",
    title: str = "手机号验证码登录成功",
) -> TestCase:
    return TestCase(
        id=tc_id,
        run_id=run.id,
        display_id=display_id,
        module=module,
        title=title,
        precondition="用户已注册",
        steps=[
            TestStep(seq=1, action="打开登录页"),
            TestStep(seq=2, action="输入手机号", data="13800138000", expected="校验通过"),
        ],
        expected="登录成功，跳转首页",
        priority=Priority.P1,
        type=TestCaseType.FUNCTIONAL,
        provenance=Provenance.LLM,
        status=status,
        generation_mode=GenerationMode.LLM,
        data_plan=[
            DataPlanItem(
                field="phone",
                strategy="example",
                value="13800138000",
                source="example",
                generator="code",
                expected_valid=True,
            )
        ],
        fingerprint=f"fp_{display_id}",
        content_hash=f"ch_{display_id}",
        remark="导出备注",
    )


def _seed(client, title: str = "退款需求") -> Run:
    uid = _session_v2_uid(client)
    run = _create_run(uid, title)
    repo.save_test_case(_case(run, tc_id=new_ulid(), display_id="TC_001", status=TestCaseStatus.REVIEWED))
    repo.save_test_case(
        _case(
            run,
            tc_id=new_ulid(),
            display_id="TC_002",
            status=TestCaseStatus.RE_REVIEW_REQUIRED,
            module="退款",
            title="退款金额超过订单金额",
        )
    )
    repo.save_test_case(
        _case(
            run,
            tc_id=new_ulid(),
            display_id="TC_003",
            status=TestCaseStatus.ARCHIVED,
            module="退款",
            title="已归档重复用例",
        )
    )
    return run


def test_excel_export_contains_active_cases_only(v2_client):
    run = _seed(v2_client)

    rv = v2_client.get(f"/api/v2/runs/{run.id}/export?format=xlsx")

    assert rv.status_code == 200
    assert rv.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in rv.headers["Content-Disposition"]
    assert ".xlsx" in rv.headers["Content-Disposition"]

    wb = load_workbook(BytesIO(rv.data))
    assert wb.sheetnames == ["测试用例", "运行信息"]
    ws = wb["测试用例"]
    assert [cell.value for cell in ws[1]] == export_lib.HEADERS

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2  # ARCHIVED 不导出
    assert {row[0] for row in rows} == {"TC_001", "TC_002"}

    login_row = next(row for row in rows if row[0] == "TC_001")
    assert login_row[1] == "登录"
    assert "1. 打开登录页" in login_row[4]
    assert "输入手机号" in login_row[4] and "13800138000" in login_row[4]
    assert login_row[6] == "phone=13800138000 (example/example)"
    assert login_row[7] == "P1"
    assert login_row[9] == "已评审"
    assert login_row[10] == "AI 生成"
    assert login_row[12] == "导出备注"

    summary = wb["运行信息"]
    values = {row[0]: row[1] for row in summary.iter_rows(min_row=2, values_only=True)}
    assert values["需求标题"] == "退款需求"
    assert values["Run ID"] == run.id
    assert values["导出用例数"] == 2
    assert values["已排除归档用例"] == 1


def test_markdown_export_contains_active_cases_only(v2_client):
    run = _seed(v2_client)

    rv = v2_client.get(f"/api/v2/runs/{run.id}/export?format=md")

    assert rv.status_code == 200
    assert rv.mimetype == "text/markdown"
    text = rv.data.decode("utf-8")
    assert "# 退款需求 - 测试用例" in text
    assert "> 用例总数: 2" in text
    assert "> 已排除归档用例: 1" in text
    assert "TC_001" in text and "TC_002" in text
    assert "TC_003" not in text
    assert "已归档重复用例" not in text
    assert "13800138000" in text
    assert "1. 打开登录页" in text
    assert "已评审" in text


def test_export_rejects_unknown_format(v2_client):
    run = _seed(v2_client)

    rv = v2_client.get(f"/api/v2/runs/{run.id}/export?format=pdf")

    assert rv.status_code == 400
    assert "不支持的导出格式" in rv.get_json()["error"]


def test_export_missing_run_returns_404(v2_client):
    assert v2_client.get(f"/api/v2/runs/{new_ulid()}/export?format=xlsx").status_code == 404


def test_export_requires_login(v2_anon):
    assert v2_anon.get(f"/api/v2/runs/{new_ulid()}/export?format=xlsx").status_code == 401


def test_export_rejects_other_users_run(v2_client):
    own_uid = _session_v2_uid(v2_client)
    _create_run(own_uid, "自己的需求")
    other_uid = new_ulid()
    repo.save_user(other_uid, "other", "", legacy_int_id=987654321)
    other_run = _create_run(other_uid, "别人的需求")
    repo.save_test_case(_case(other_run, tc_id=new_ulid(), display_id="TC_OTHER", status=TestCaseStatus.REVIEWED))

    rv = v2_client.get(f"/api/v2/runs/{other_run.id}/export?format=xlsx")

    assert rv.status_code == 404  # 不暴露他人 Run 是否存在


def test_export_gating_503_when_not_ready(v2_client, monkeypatch):
    import web

    run = _seed(v2_client)
    monkeypatch.setitem(web.app.config, "V2_READY", False)

    rv = v2_client.get(f"/api/v2/runs/{run.id}/export?format=xlsx")

    assert rv.status_code == 503


def test_export_empty_run(v2_client):
    uid = _session_v2_uid(v2_client)
    run = _create_run(uid, "空需求")

    rv = v2_client.get(f"/api/v2/runs/{run.id}/export?format=xlsx")

    assert rv.status_code == 200
    ws = load_workbook(BytesIO(rv.data))["测试用例"]
    assert len(list(ws.iter_rows(min_row=2, values_only=True))) == 0


def test_safe_export_filename_strips_path_chars():
    name = export_lib.safe_export_filename('退款/测试: "加急"?', ext="xlsx", exported_at=datetime(2026, 1, 2, 3, 4, 5))
    assert "/" not in name and ":" not in name and '"' not in name and "?" not in name
    assert name.endswith("_测试用例_20260102_030405.xlsx")
